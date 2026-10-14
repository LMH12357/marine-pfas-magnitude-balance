"""Sensitivity analysis of left-censored PFAS data treatment and GDR model performance."""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import operator
import os
import sys
import importlib
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.base import clone
from sklearn.ensemble import RandomForestRegressor
from sklearn.model_selection import (
    GridSearchCV,
    GroupKFold,
    KFold,
    LeaveOneGroupOut,
    train_test_split,
)
from xgboost import XGBRegressor

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pfas_schema import (
    COVARIATES,
    MODEL_GROUPS,
    normalize_feature_columns,
    schema_digest,
)
from metadata_schema import normalize_metadata

SEED = 42
N_OUTER_FOLDS = 5
N_INNER_FOLDS = 3
N_JOBS = 1
ND_MDL = "ND/<MDL"
BELOW_MQL = "<MQL"
MEASURED = "Measured"
UNMEASURED = "Unmeasured/unreported"


@dataclass(frozen=True)
class GDRDefinition:
    numerator: tuple[str, ...]
    denominator: tuple[str, ...]


GDR_DEFINITIONS = {
    "Ratio_Short_Long_Ocean": GDRDefinition(
        MODEL_GROUPS["Short_PFSA"] + MODEL_GROUPS["Short_PFCA"],
        MODEL_GROUPS["Long_PFSA"] + MODEL_GROUPS["Long_PFCA"],
    ),
    "Ratio_Long_PFSA_PFCA_Ocean": GDRDefinition(
        MODEL_GROUPS["Long_PFSA"],
        MODEL_GROUPS["Long_PFCA"],
    ),
    "Ratio_Short_PFSA_PFCA_Ocean": GDRDefinition(
        MODEL_GROUPS["Short_PFSA"],
        MODEL_GROUPS["Short_PFCA"],
    ),
}

PFAS_COLUMNS = sorted(
    {
        column
        for definition in GDR_DEFINITIONS.values()
        for column in definition.numerator + definition.denominator
    }
)


@dataclass(frozen=True)
class CensoringScheme:
    nd_mdl_fraction: float
    mql_fraction: float
    role: str


SCHEMES = {
    "S0_submitted": CensoringScheme(1 / np.sqrt(2), 1 / 2, "Primary"),
    "S1_half_limit": CensoringScheme(1 / 2, 1 / 2, "Sensitivity"),
    "S2_sqrt2_limit": CensoringScheme(1 / np.sqrt(2), 1 / np.sqrt(2), "Sensitivity"),
    "S3_lower_bound_zero": CensoringScheme(0.0, 0.0, "Theoretical lower bound"),
    "S4_upper_bound_limit": CensoringScheme(1.0, 1.0, "Theoretical upper bound"),
}

RF_GRID = {
    "n_estimators": [100, 200],
    "max_depth": [3, 5, 8],
    "max_features": [0.6, 0.8, 1.0],
    "max_samples": [0.6, 0.8, 1.0],
}

XGB_GRID = {
    "n_estimators": [100, 200],
    "max_depth": [3, 5],
    "learning_rate": [0.01, 0.1],
    "subsample": [0.6, 0.8],
    "colsample_bytree": [0.6, 0.8],
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def evaluate_numeric_expression(expression: str) -> float:
    tree = ast.parse(expression, mode="eval")
    binary = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.Pow: operator.pow,
    }
    unary = {ast.UAdd: operator.pos, ast.USub: operator.neg}

    def visit(node: ast.AST) -> float:
        if isinstance(node, ast.Expression):
            return visit(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.BinOp) and type(node.op) in binary:
            return float(binary[type(node.op)](visit(node.left), visit(node.right)))
        if isinstance(node, ast.UnaryOp) and type(node.op) in unary:
            return float(unary[type(node.op)](visit(node.operand)))
        raise ValueError(f"Unsupported formula: {expression}")

    return visit(tree)


def parse_pfas_cell(value: object, data_type: str) -> tuple[str, float | None]:
    if value is None:
        return UNMEASURED, None
    if data_type != "f":
        number = float(value)
        if not np.isfinite(number) or number < 0:
            raise ValueError("PFAS concentrations must be finite and nonnegative")
        return MEASURED, number
    formula = str(value).replace(" ", "")
    body = formula[1:] if formula.startswith("=") else formula
    lowered = body.lower()
    if lowered.endswith("/2^0.5"):
        return ND_MDL, evaluate_numeric_expression(body[:-6])
    if lowered.endswith("/sqrt(2)"):
        return ND_MDL, evaluate_numeric_expression(body[:-8])
    if lowered.endswith("/2"):
        return BELOW_MQL, evaluate_numeric_expression(body[:-2])
    raise ValueError(f"Unrecognized PFAS formula: {formula}")


def reconstruct_pfas_values(workbook_path: Path, n_rows: int):
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = workbook.active
    headers = {
        cell.value: cell.column for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    }
    available = [column for column in PFAS_COLUMNS if column in headers]
    if not available:
        workbook.close()
        raise ValueError("No model-panel PFAS columns occur in the workbook")
    index_to_column = {headers[column] - 1: column for column in available}
    first_index = min(index_to_column)
    last_index = max(index_to_column)
    values = {scheme: {column: [] for column in available} for scheme in SCHEMES}
    statuses = {column: [] for column in available}
    limits = {column: [] for column in available}
    formulas = {column: [] for column in available}
    try:
        rows = sheet.iter_rows(
            min_row=2,
            max_row=n_rows + 1,
            min_col=first_index + 1,
            max_col=last_index + 1,
        )
        for row in rows:
            for local_index, cell in enumerate(row, start=first_index):
                column = index_to_column.get(local_index)
                if column is None:
                    continue
                try:
                    status, value_or_limit = parse_pfas_cell(cell.value, cell.data_type)
                except (ValueError, TypeError) as error:
                    raise ValueError(
                        f"{cell.coordinate} ({column}): {error}"
                    ) from error
                if value_or_limit is not None and (
                    not np.isfinite(value_or_limit) or value_or_limit < 0
                ):
                    raise ValueError(
                        f"{cell.coordinate}: invalid analytical limit or concentration"
                    )
                statuses[column].append(status)
                limits[column].append(
                    value_or_limit if status in {ND_MDL, BELOW_MQL} else np.nan
                )
                formulas[column].append(
                    str(cell.value) if cell.data_type == "f" else ""
                )
                for scheme_name, scheme in SCHEMES.items():
                    if status == ND_MDL:
                        reconstructed = float(value_or_limit) * scheme.nd_mdl_fraction
                    elif status == BELOW_MQL:
                        reconstructed = float(value_or_limit) * scheme.mql_fraction
                    elif status == MEASURED:
                        reconstructed = float(value_or_limit)
                    else:
                        reconstructed = np.nan
                    values[scheme_name][column].append(reconstructed)
    finally:
        workbook.close()
    frames = {name: pd.DataFrame(frame) for name, frame in values.items()}
    return frames, pd.DataFrame(statuses), pd.DataFrame(limits), pd.DataFrame(formulas)


def inspect_censoring_evidence(workbook_path: Path, n_rows: int):
    """Audit formula provenance and flag numeric red cells with unknown limits.

    Red font is a provenance cue, not enough information to infer a limit or
    divisor. Such cells must not silently enter sensitivity runs as measured.
    """
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = workbook.active
    headers = {
        cell.value: cell.column - 1
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    }
    available = {index: name for name, index in headers.items() if name in PFAS_COLUMNS}
    counts, formula_rows, unresolved = {}, {}, []
    try:
        for row in sheet.iter_rows(min_row=2, max_row=n_rows + 1):
            for index, analyte in available.items():
                cell = row[index]
                color = cell.font.color if cell.font is not None else None
                rgb = (
                    str(color.rgb).upper()
                    if color is not None and color.type == "rgb"
                    else ""
                )
                red = rgb.endswith("FF0000")
                kind = (
                    "formula"
                    if cell.data_type == "f"
                    else "blank" if cell.value is None else "value"
                )
                key = f"{'red' if red else 'not_red'}_{kind}"
                counts[key] = counts.get(key, 0) + 1
                issue = None
                if red and kind == "value":
                    issue = (
                        "Red numeric or text cell has no recoverable censoring formula"
                    )
                try:
                    status, limit = parse_pfas_cell(cell.value, cell.data_type)
                    if limit is not None and (not np.isfinite(limit) or limit < 0):
                        issue = "Invalid concentration or analytical limit"
                    if kind == "formula":
                        expression = str(cell.value)
                        if expression not in formula_rows:
                            formula_rows[expression] = {
                                "formula": expression,
                                "status": status,
                                "analytical_limit": limit,
                                "cell_count": 0,
                            }
                        formula_rows[expression]["cell_count"] += 1
                except (ValueError, TypeError, SyntaxError) as error:
                    issue = str(error)
                if issue:
                    unresolved.append(
                        {
                            "cell": cell.coordinate,
                            "analyte": analyte,
                            "value": str(cell.value),
                            "font_rgb": rgb,
                            "issue": issue,
                        }
                    )
    finally:
        workbook.close()
    columns = ["cell", "analyte", "value", "font_rgb", "issue"]
    return (
        counts,
        pd.DataFrame(formula_rows.values()),
        pd.DataFrame(unresolved, columns=columns),
    )


def reconstruction_audit(cached_data, reconstructed, statuses, limits, formulas):
    rows = []
    primary = reconstructed["S0_submitted"]
    for analyte in primary.columns:
        for row_index in primary.index:
            workbook_value = cached_data.at[row_index, analyte]
            primary_value = primary.at[row_index, analyte]
            if pd.isna(workbook_value) and pd.isna(primary_value):
                difference = np.nan
                matches = True
            else:
                difference = float(primary_value) - float(workbook_value)
                matches = bool(
                    np.isclose(primary_value, workbook_value, rtol=1e-12, atol=1e-12)
                )
            rows.append(
                {
                    "row_index": row_index,
                    "excel_row": row_index + 2,
                    "analyte": analyte,
                    "source_formula": formulas.at[row_index, analyte],
                    "status": statuses.at[row_index, analyte],
                    "analytical_limit": limits.at[row_index, analyte],
                    "workbook_cached_value": workbook_value,
                    "S0_value": primary_value,
                    "S1_value": reconstructed["S1_half_limit"].at[row_index, analyte],
                    "S2_value": reconstructed["S2_sqrt2_limit"].at[row_index, analyte],
                    "S3_value": reconstructed["S3_lower_bound_zero"].at[
                        row_index, analyte
                    ],
                    "S4_value": reconstructed["S4_upper_bound_limit"].at[
                        row_index, analyte
                    ],
                    "S0_minus_workbook": difference,
                    "S0_matches_workbook": matches,
                }
            )
    audit = pd.DataFrame(rows)
    if not audit["S0_matches_workbook"].all():
        raise ValueError("S0 does not reproduce all cached PFAS values in DATA.xlsx.")
    return audit


def target_reproduction_audit(cached_data, primary_targets):
    workbook_targets, _ = construct_gdrs(cached_data)
    rows = []
    for target in GDR_DEFINITIONS:
        difference = primary_targets[target] - workbook_targets[target]
        finite = difference[np.isfinite(difference)]
        maximum = float(finite.abs().max()) if len(finite) else 0.0
        matches = bool(
            np.allclose(
                primary_targets[target],
                workbook_targets[target],
                equal_nan=True,
                rtol=1e-12,
                atol=1e-12,
            )
        )
        rows.append(
            {
                "target": target,
                "rows": len(cached_data),
                "maximum_absolute_difference": maximum,
                "S0_matches_submitted_GDR": matches,
            }
        )
    audit = pd.DataFrame(rows)
    if not audit["S0_matches_submitted_GDR"].all():
        raise ValueError("S0 does not reproduce the submitted GDR values.")
    return audit


def resolve_column_names(frame: pd.DataFrame) -> pd.DataFrame:
    aliases = {"Temperature (��C)": "Temperature (°C)", "Salinity (��)": "Salinity (‰)"}
    return normalize_feature_columns(
        frame.rename(
            columns={
                column: aliases[column] for column in frame.columns if column in aliases
            }
        )
    )


def construct_gdrs(pfas_values: pd.DataFrame):
    values = pfas_values[
        [column for column in PFAS_COLUMNS if column in pfas_values]
    ].apply(pd.to_numeric, errors="raise")
    if np.isinf(values.to_numpy(dtype=float)).any() or (values < 0).any().any():
        raise ValueError("PFAS concentrations cannot be negative or infinite")
    targets = pd.DataFrame(index=pfas_values.index)
    audit_rows = []
    for target, definition in GDR_DEFINITIONS.items():
        numerator_columns = [
            column for column in definition.numerator if column in pfas_values.columns
        ]
        denominator_columns = [
            column for column in definition.denominator if column in pfas_values.columns
        ]
        numerator_count = pfas_values[numerator_columns].notna().sum(axis=1)
        denominator_count = pfas_values[denominator_columns].notna().sum(axis=1)
        numerator = pfas_values[numerator_columns].sum(axis=1, min_count=1)
        denominator = pfas_values[denominator_columns].sum(axis=1, min_count=1)
        targets[target] = numerator / denominator.where(denominator > 0)
        audit_rows.append(
            pd.DataFrame(
                {
                    "row_index": pfas_values.index,
                    "target": target,
                    "numerator_component_count": numerator_count,
                    "denominator_component_count": denominator_count,
                    "numerator_sum": numerator,
                    "denominator_sum": denominator,
                    "ratio_available": targets[target].notna(),
                }
            )
        )
    return targets, pd.concat(audit_rows, ignore_index=True)


def calculate_censoring_contribution(values: pd.DataFrame, statuses: pd.DataFrame):
    rows = []
    for target, definition in GDR_DEFINITIONS.items():
        for component, declared_columns in (
            ("Numerator", definition.numerator),
            ("Denominator", definition.denominator),
        ):
            columns = [
                column for column in declared_columns if column in values.columns
            ]
            censored = statuses[columns].isin([ND_MDL, BELOW_MQL])
            total_mass = values[columns].sum(axis=1, min_count=1)
            censored_mass = (
                values[columns].where(censored).sum(axis=1, min_count=1).fillna(0)
            )
            observed_count = values[columns].notna().sum(axis=1)
            censored_count = censored.sum(axis=1)
            rows.append(
                pd.DataFrame(
                    {
                        "row_index": values.index,
                        "target": target,
                        "component": component,
                        "total_mass": total_mass,
                        "censored_mass": censored_mass,
                        "censored_mass_fraction": censored_mass
                        / total_mass.where(total_mass > 0),
                        "observed_component_count": observed_count,
                        "censored_component_count": censored_count,
                        "censored_component_fraction": censored_count
                        / observed_count.where(observed_count > 0),
                    }
                )
            )
    return pd.concat(rows, ignore_index=True)


def regression_metrics(observed: np.ndarray, predicted: np.ndarray):
    residual = observed - predicted
    mse = float(np.mean(residual**2))
    rmse = float(np.sqrt(mse))
    mae = float(np.mean(np.abs(residual)))
    total_sum_squares = float(np.sum((observed - observed.mean()) ** 2))
    r2 = (
        1 - float(np.sum(residual**2)) / total_sum_squares
        if total_sum_squares > 0
        else np.nan
    )
    return {"R2": r2, "MSE": mse, "RMSE": rmse, "MAE": mae}


def model_definitions():
    return {
        "RF": (
            RandomForestRegressor(
                random_state=SEED, criterion="absolute_error", n_jobs=1
            ),
            RF_GRID,
        ),
        "XGB": (
            XGBRegressor(random_state=SEED, objective="reg:pseudohubererror", n_jobs=1),
            XGB_GRID,
        ),
    }


def fit_original_protocol(features, targets_by_scheme, scheme_names):
    performance_rows = []
    prediction_rows = []
    champion_rows = []
    for scheme_name in scheme_names:
        for target in GDR_DEFINITIONS:
            frame = features.copy()
            frame[target] = targets_by_scheme[scheme_name][target]
            frame = frame.dropna(subset=[target] + COVARIATES)
            x = frame[COVARIATES]
            y_unclipped = frame[target]
            lower = float(y_unclipped.quantile(0.025))
            upper = float(y_unclipped.quantile(0.975))
            y = y_unclipped.clip(lower, upper)
            x_train, x_test, y_train, y_test = train_test_split(
                x, y, test_size=0.20, random_state=SEED
            )
            target_results = []
            for model_name, (estimator, parameter_grid) in model_definitions().items():
                search = GridSearchCV(
                    estimator, parameter_grid, cv=5, scoring="r2", n_jobs=N_JOBS
                )
                search.fit(x_train, y_train)
                prediction = search.best_estimator_.predict(x_test)
                metrics = regression_metrics(y_test.to_numpy(dtype=float), prediction)
                result = {
                    "evaluation_track": "original_protocol",
                    "target": target,
                    "scheme": scheme_name,
                    "model": model_name,
                    "n_total": len(frame),
                    "train_n": len(x_train),
                    "test_n": len(x_test),
                    "global_clip_lower": lower,
                    "global_clip_upper": upper,
                    "grid_scoring": "r2",
                    "best_parameters": json.dumps(search.best_params_, sort_keys=True),
                    **metrics,
                }
                performance_rows.append(result)
                target_results.append(result)
                for row_index, observed, predicted in zip(
                    y_test.index, y_test, prediction
                ):
                    prediction_rows.append(
                        {
                            "evaluation_track": "original_protocol",
                            "row_index": row_index,
                            "target": target,
                            "scheme": scheme_name,
                            "model": model_name,
                            "observed": observed,
                            "predicted": predicted,
                        }
                    )
                print(f"Original protocol | {target} | {scheme_name} | {model_name}")
            champion = max(target_results, key=lambda row: row["R2"])
            champion_rows.append(
                {
                    "target": target,
                    "scheme": scheme_name,
                    "architecture": champion["model"],
                    "R2": champion["R2"],
                    "MSE": champion["MSE"],
                    "RMSE": champion["RMSE"],
                    "MAE": champion["MAE"],
                }
            )
    return (
        pd.DataFrame(performance_rows),
        pd.DataFrame(prediction_rows),
        pd.DataFrame(champion_rows),
    )


def fit_paired_fixed_sensitivity(features, targets_by_scheme, scheme_names):
    performance_rows = []
    prediction_rows = []
    for target in GDR_DEFINITIONS:
        common = features.index[features[COVARIATES].notna().all(axis=1)]
        for scheme_name in scheme_names:
            values = targets_by_scheme[scheme_name][target]
            common = common.intersection(
                values.index[values.notna() & np.isfinite(values)]
            )
        train_index, test_index = train_test_split(
            common.to_numpy(), test_size=0.20, random_state=SEED
        )
        x_train = features.loc[train_index, COVARIATES]
        x_test = features.loc[test_index, COVARIATES]
        primary = targets_by_scheme["S0_submitted"].loc[common, target]
        primary_lower = float(primary.quantile(0.025))
        primary_upper = float(primary.quantile(0.975))
        primary_training = primary.clip(primary_lower, primary_upper).loc[train_index]

        for model_name, (estimator, parameter_grid) in model_definitions().items():
            search = GridSearchCV(
                estimator, parameter_grid, cv=5, scoring="r2", n_jobs=N_JOBS
            )
            search.fit(x_train, primary_training)
            fixed_parameters = search.best_params_
            for scheme_name in scheme_names:
                values = targets_by_scheme[scheme_name].loc[common, target]
                lower = float(values.quantile(0.025))
                upper = float(values.quantile(0.975))
                clipped = values.clip(lower, upper)
                y_train = clipped.loc[train_index]
                y_test = clipped.loc[test_index]
                fitted = clone(estimator).set_params(**fixed_parameters)
                fitted.fit(x_train, y_train)
                prediction = fitted.predict(x_test)
                performance_rows.append(
                    {
                        "evaluation_track": "paired_fixed_original_split",
                        "target": target,
                        "scheme": scheme_name,
                        "model": model_name,
                        "paired_n": len(common),
                        "train_n": len(train_index),
                        "test_n": len(test_index),
                        "global_clip_lower": lower,
                        "global_clip_upper": upper,
                        "hyperparameter_source": "S0_submitted",
                        "best_parameters": json.dumps(fixed_parameters, sort_keys=True),
                        **regression_metrics(y_test.to_numpy(dtype=float), prediction),
                    }
                )
                for row_index, observed, predicted in zip(
                    test_index, y_test, prediction
                ):
                    prediction_rows.append(
                        {
                            "evaluation_track": "paired_fixed_original_split",
                            "row_index": row_index,
                            "target": target,
                            "scheme": scheme_name,
                            "model": model_name,
                            "observed": observed,
                            "predicted": predicted,
                        }
                    )
                print(f"Paired fixed | {target} | {scheme_name} | {model_name}")
    return pd.DataFrame(performance_rows), pd.DataFrame(prediction_rows)


def single_split_deltas(performance):
    metrics = ["R2", "MSE", "RMSE", "MAE"]
    primary = performance[performance["scheme"] == "S0_submitted"].set_index(
        ["target", "model"]
    )
    rows = []
    for _, row in performance.iterrows():
        baseline = primary.loc[(row["target"], row["model"])]
        rows.append(
            {
                "evaluation_track": row["evaluation_track"],
                "target": row["target"],
                "scheme": row["scheme"],
                "model": row["model"],
                **{
                    f"Delta_{metric}": row[metric] - baseline[metric]
                    for metric in metrics
                },
            }
        )
    return pd.DataFrame(rows)


def compare_historical_report(performance, report_path):
    report = pd.read_csv(report_path)
    current = performance[performance["scheme"] == "S0_submitted"]
    rows = []
    for _, reference in report.iterrows():
        target = reference["Target"]
        for model in ["RF", "XGB"]:
            result = current[
                (current["target"] == target) & (current["model"] == model)
            ].iloc[0]
            row = {"target": target, "model": model}
            matches = True
            for metric in ["R2", "RMSE", "MAE"]:
                reference_value = float(reference[f"{model}_{metric}"])
                current_value = float(result[metric])
                metric_matches = round(current_value, 4) == round(reference_value, 4)
                row[f"historical_{metric}"] = reference_value
                row[f"reproduced_{metric}"] = current_value
                row[f"difference_{metric}"] = current_value - reference_value
                row[f"matches_{metric}_at_4_decimals"] = metric_matches
                matches = matches and metric_matches
            row["all_metrics_match_at_4_decimals"] = matches
            rows.append(row)
    audit = pd.DataFrame(rows)
    if not audit["all_metrics_match_at_4_decimals"].all():
        raise ValueError(
            "S0 does not reproduce the historical model report at four decimal places."
        )
    return audit


def outer_splits(frame: pd.DataFrame, validation: str):
    if validation == "sampling_group":
        groups = frame["Sampling_Group"].astype(str)
        return list(
            GroupKFold(n_splits=min(N_OUTER_FOLDS, groups.nunique())).split(
                frame, groups=groups
            )
        )
    if validation == "spatial_block":
        longitude_group = pd.cut(frame["Longitude"], bins=3, labels=False)
        latitude_group = pd.cut(frame["Latitude"], bins=3, labels=False)
        groups = longitude_group.astype(str) + "_" + latitude_group.astype(str)
        return list(
            GroupKFold(n_splits=min(N_OUTER_FOLDS, groups.nunique())).split(
                frame, groups=groups
            )
        )
    if validation == "leave_one_sea_region_out":
        groups = frame["Sea_Region"].astype(str)
        return list(LeaveOneGroupOut().split(frame, groups=groups))
    return list(
        KFold(n_splits=N_OUTER_FOLDS, shuffle=True, random_state=SEED).split(frame)
    )


def fit_sensitivity_validation(features, targets_by_scheme, scheme_names, validation):
    performance_rows = []
    prediction_rows = []
    for target in GDR_DEFINITIONS:
        common = features.index[features[COVARIATES].notna().all(axis=1)]
        for scheme_name in scheme_names:
            scheme_targets = targets_by_scheme[scheme_name]
            valid = scheme_targets.index[
                scheme_targets[target].notna() & np.isfinite(scheme_targets[target])
            ]
            common = common.intersection(valid)
        analysis = features.loc[common]
        splits = outer_splits(analysis, validation)

        for fold_number, (train_position, test_position) in enumerate(splits, start=1):
            train_index = analysis.index[train_position]
            test_index = analysis.index[test_position]
            x_train = analysis.loc[train_index, COVARIATES]
            x_test = analysis.loc[test_index, COVARIATES]
            inner_groups = analysis.loc[train_index, "Sampling_Group"].astype(str)
            inner_cv = GroupKFold(n_splits=min(N_INNER_FOLDS, inner_groups.nunique()))
            primary_train = (
                targets_by_scheme["S0_submitted"]
                .loc[train_index, target]
                .to_numpy(dtype=float)
            )
            primary_lower = float(np.quantile(primary_train, 0.025))
            primary_upper = float(np.quantile(primary_train, 0.975))
            primary_training_target = np.clip(
                primary_train, primary_lower, primary_upper
            )
            model_definitions = {
                "RF": (
                    RandomForestRegressor(
                        random_state=SEED, criterion="absolute_error", n_jobs=1
                    ),
                    RF_GRID,
                ),
                "XGB": (
                    XGBRegressor(
                        random_state=SEED, objective="reg:pseudohubererror", n_jobs=1
                    ),
                    XGB_GRID,
                ),
            }

            for model_name, (estimator, parameter_grid) in model_definitions.items():
                search = GridSearchCV(
                    estimator,
                    parameter_grid,
                    scoring="neg_mean_squared_error",
                    cv=inner_cv,
                    n_jobs=N_JOBS,
                )
                search.fit(x_train, primary_training_target, groups=inner_groups)
                fixed_parameters = search.best_params_

                for scheme_name in scheme_names:
                    target_frame = targets_by_scheme[scheme_name]
                    y_train_observed = target_frame.loc[train_index, target].to_numpy(
                        dtype=float
                    )
                    y_test = target_frame.loc[test_index, target].to_numpy(dtype=float)
                    lower = float(np.quantile(y_train_observed, 0.025))
                    upper = float(np.quantile(y_train_observed, 0.975))
                    y_train = np.clip(y_train_observed, lower, upper)
                    fitted_model = clone(estimator).set_params(**fixed_parameters)
                    fitted_model.fit(x_train, y_train)
                    prediction = fitted_model.predict(x_test)
                    performance_rows.append(
                        {
                            "evaluation_track": "sensitivity_grouped_nested",
                            "target": target,
                            "scheme": scheme_name,
                            "scheme_role": SCHEMES[scheme_name].role,
                            "model": model_name,
                            "validation": validation,
                            "fold": fold_number,
                            "train_n": len(train_index),
                            "test_n": len(test_index),
                            "training_clip_lower": lower,
                            "training_clip_upper": upper,
                            "hyperparameter_source": "S0_submitted",
                            "best_parameters": json.dumps(
                                fixed_parameters, sort_keys=True
                            ),
                            **regression_metrics(y_test, prediction),
                        }
                    )
                    for row_index, observed, predicted in zip(
                        test_index, y_test, prediction
                    ):
                        prediction_rows.append(
                            {
                                "evaluation_track": "sensitivity_grouped_nested",
                                "row_index": row_index,
                                "target": target,
                                "scheme": scheme_name,
                                "model": model_name,
                                "validation": validation,
                                "fold": fold_number,
                                "observed": observed,
                                "predicted": predicted,
                            }
                        )
                    print(
                        f"{target} | {scheme_name} | {model_name} | fold {fold_number}/{len(splits)}"
                    )
    return pd.DataFrame(performance_rows), pd.DataFrame(prediction_rows)


def summarize_performance(per_fold: pd.DataFrame):
    metric_columns = ["R2", "MSE", "RMSE", "MAE"]
    grouping = ["target", "scheme", "scheme_role", "model", "validation"]
    summary = per_fold.groupby(grouping)[metric_columns].agg(["mean", "std", "median"])
    summary.columns = [f"{metric}_{statistic}" for metric, statistic in summary.columns]
    summary = summary.reset_index()
    primary = summary[summary["scheme"] == "S0_submitted"].set_index(
        ["target", "model", "validation"]
    )
    delta_rows = []
    for _, row in summary.iterrows():
        baseline = primary.loc[(row["target"], row["model"], row["validation"])]
        delta_rows.append(
            {
                "target": row["target"],
                "scheme": row["scheme"],
                "model": row["model"],
                "validation": row["validation"],
                **{
                    f"Delta_{metric}_mean": row[f"{metric}_mean"]
                    - baseline[f"{metric}_mean"]
                    for metric in metric_columns
                },
            }
        )
    return summary, pd.DataFrame(delta_rows)


def summarize_pooled_oof_performance(predictions: pd.DataFrame) -> pd.DataFrame:
    """Compute pooled held-out metrics on the untransformed group-ratio scale."""
    grouping = ["target", "scheme", "model", "validation"]
    missing = sorted(
        set(grouping + ["observed", "predicted"]).difference(predictions.columns)
    )
    if missing:
        raise ValueError(f"Missing pooled OOF columns: {missing}")

    rows = []
    for keys, group in predictions.groupby(grouping, sort=True):
        observed = group["observed"].to_numpy(dtype=float)
        predicted = group["predicted"].to_numpy(dtype=float)
        if not np.isfinite(observed).all() or not np.isfinite(predicted).all():
            raise ValueError(f"Non-finite pooled OOF values for {keys}")
        rows.append(
            {
                **dict(zip(grouping, keys)),
                "N": len(group),
                **regression_metrics(observed, predicted),
            }
        )

    pooled = pd.DataFrame(rows, columns=grouping + ["N", "R2", "MSE", "RMSE", "MAE"])
    if pooled.empty:
        return pooled.assign(Delta_R2=np.nan, Delta_RMSE=np.nan)
    reference = pooled[pooled["scheme"] == "S0_submitted"].set_index(
        ["target", "model", "validation"]
    )
    for metric in ("R2", "RMSE"):
        pooled[f"Delta_{metric}"] = [
            row[metric]
            - reference.loc[(row["target"], row["model"], row["validation"]), metric]
            for _, row in pooled.iterrows()
        ]
    return pooled


def ratio_stability(targets_by_scheme: dict[str, pd.DataFrame]):
    rows = []
    primary = targets_by_scheme["S0_submitted"]
    for target in GDR_DEFINITIONS:
        common = primary.index[primary[target].notna() & np.isfinite(primary[target])]
        for target_frame in targets_by_scheme.values():
            valid = target_frame.index[
                target_frame[target].notna() & np.isfinite(target_frame[target])
            ]
            common = common.intersection(valid)
        baseline = primary.loc[common, target]
        for scheme_name, target_frame in targets_by_scheme.items():
            values = target_frame.loc[common, target]
            relative_change = (values - baseline) / baseline.replace(0, np.nan)
            rows.append(
                {
                    "target": target,
                    "scheme": scheme_name,
                    "n": len(common),
                    "median": values.median(),
                    "p2.5": values.quantile(0.025),
                    "p97.5": values.quantile(0.975),
                    "Spearman_vs_submitted": values.rank().corr(baseline.rank()),
                    "median_absolute_relative_change": relative_change.abs().median(),
                }
            )
    return pd.DataFrame(rows)


def parse_arguments(argv=None):
    project_root = (
        Path(os.environ.get("PFAS_PROJECT_ROOT", Path(__file__).resolve().parents[2]))
        .expanduser()
        .resolve()
    )
    data_root = (
        Path(os.environ.get("PFAS_DATA_ROOT", project_root / "data"))
        .expanduser()
        .resolve()
    )
    output_root = (
        Path(os.environ.get("PFAS_OUTPUT_ROOT", project_root / "outputs"))
        .expanduser()
        .resolve()
    )
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", type=Path, default=data_root / "raw" / "DATA.xlsx")
    parser.add_argument(
        "--source-load",
        type=Path,
        default=None,
        help="Current source-preparation Step4 CSV; required when routing features from raw DATA.xlsx.",
    )
    parser.add_argument(
        "--boundary",
        type=Path,
        default=data_root / "reference" / "geospatial" / "china_coastline_boundary.shp",
    )
    parser.add_argument("--feature-table", type=Path, default=None)
    parser.add_argument("--routing-directory", type=Path, default=None)
    parser.add_argument(
        "--reference-report",
        type=Path,
        default=None,
        help="Optional report for an explicitly requested historical reproduction; not used automatically.",
    )
    parser.add_argument(
        "--output-dir", type=Path, default=output_root / "04_preprocessing_sensitivity"
    )
    parser.add_argument(
        "--analysis-mode",
        choices=["all", "legacy", "validation", "prepare"],
        default="all",
    )
    parser.add_argument(
        "--schemes", nargs="+", choices=list(SCHEMES), default=list(SCHEMES)
    )
    parser.add_argument(
        "--validation",
        choices=[
            "sampling_group",
            "spatial_block",
            "leave_one_sea_region_out",
            "random",
        ],
        default="sampling_group",
    )
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument(
        "--n-jobs",
        type=int,
        default=1,
        help="Parallel grid-search jobs; estimators use one thread to avoid nested oversubscription.",
    )
    arguments = parser.parse_args(argv)
    if "S0_submitted" not in arguments.schemes:
        parser.error("S0_submitted is required as the paired reference scheme")
    if arguments.n_jobs == 0:
        parser.error("--n-jobs cannot be zero")
    if (
        not arguments.prepare_only
        and arguments.analysis_mode != "prepare"
        and arguments.feature_table is None
        and arguments.source_load is None
    ):
        parser.error("Provide current --source-load or an explicit --feature-table")
    return arguments


def validate_feature_alignment(data, features):
    """Reject stale or reordered raw inputs before pairing censoring alternatives."""
    expected = np.arange(len(data), dtype=int)
    if "row_index" not in features or not np.array_equal(
        features["row_index"].to_numpy(), expected
    ):
        raise ValueError(
            "Feature rows must preserve every DATA.xlsx row in its original order"
        )
    for column in ["Longitude", "Latitude"] + [
        name for name in PFAS_COLUMNS if name in data
    ]:
        if column not in features or not np.allclose(
            pd.to_numeric(data[column]).to_numpy(dtype=float),
            pd.to_numeric(features[column]).to_numpy(dtype=float),
            rtol=1e-12,
            atol=1e-12,
            equal_nan=True,
        ):
            raise ValueError(f"Feature table disagrees with DATA.xlsx: {column}")


def main():
    global N_JOBS
    arguments = parse_arguments()
    N_JOBS = arguments.n_jobs
    if arguments.prepare_only:
        arguments.analysis_mode = "prepare"
    primary = importlib.import_module("01_primary_model_training")
    input_provenance = {
        "data": {
            "path": str(arguments.data.resolve()),
            "sha256": primary.sha256(arguments.data),
        }
    }
    if arguments.analysis_mode != "prepare":
        input_provenance.update(primary.training_input_provenance(arguments))
    primary.prepare_output_directory(arguments.output_dir, input_provenance)
    data = resolve_column_names(pd.read_excel(arguments.data))
    censoring_counts, formula_syntax, unresolved = inspect_censoring_evidence(
        arguments.data, len(data)
    )
    formula_syntax.to_csv(
        arguments.output_dir / "Censoring_Formula_Syntax.csv", index=False
    )
    unresolved.to_csv(
        arguments.output_dir / "Censoring_Unresolved_Cells.csv", index=False
    )
    if not unresolved.empty:
        raise ValueError(
            f"{len(unresolved)} unresolved PFAS provenance cells; inspect Censoring_Unresolved_Cells.csv"
        )
    reconstructed, statuses, limits, formulas = reconstruct_pfas_values(
        arguments.data, len(data)
    )
    targets_by_scheme = {}
    sample_flow = []
    contributions = []

    for scheme_name, values in reconstructed.items():
        targets, audit = construct_gdrs(values)
        targets_by_scheme[scheme_name] = targets
        audit.insert(0, "scheme", scheme_name)
        sample_flow.append(audit)
        contribution = calculate_censoring_contribution(values, statuses)
        contribution.insert(1, "scheme", scheme_name)
        contributions.append(contribution)

    cell_audit = reconstruction_audit(data, reconstructed, statuses, limits, formulas)
    target_audit = target_reproduction_audit(data, targets_by_scheme["S0_submitted"])
    submitted_targets, _ = construct_gdrs(data)
    targets_by_scheme["S0_submitted"] = submitted_targets
    cell_audit.to_csv(
        arguments.output_dir / "S0_Cell_Reconstruction_Audit.csv", index=False
    )
    target_audit.to_csv(
        arguments.output_dir / "S0_GDR_Reproduction_Audit.csv", index=False
    )
    pd.concat(targets_by_scheme, axis=1).to_csv(
        arguments.output_dir / "GDR_Targets_by_Scheme.csv", index_label="row_index"
    )
    pd.concat(sample_flow, ignore_index=True).to_csv(
        arguments.output_dir / "GDR_Sample_Flow.csv", index=False
    )
    pd.concat(contributions, ignore_index=True).to_csv(
        arguments.output_dir / "GDR_Censoring_Contribution.csv", index=False
    )
    statuses.reset_index(names="row_index").melt(
        id_vars="row_index", var_name="analyte", value_name="status"
    ).to_csv(arguments.output_dir / "PFAS_Cell_Status.csv", index=False)
    limits.reset_index(names="row_index").melt(
        id_vars="row_index", var_name="analyte", value_name="analytical_limit"
    ).to_csv(arguments.output_dir / "PFAS_Analytical_Limits.csv", index=False)
    ratio_stability(targets_by_scheme).to_csv(
        arguments.output_dir / "GDR_Ratio_Stability.csv", index=False
    )

    manifest = {
        "analysis_schema_sha256": schema_digest(),
        "pfas_schema_sha256": schema_digest(),
        "code_sha256": file_sha256(Path(__file__)),
        "runtime_versions": primary.runtime_versions(),
        "input_provenance": input_provenance,
        "data": str(arguments.data.resolve()),
        "data_sha256": file_sha256(arguments.data),
        "rows": len(data),
        "censoring_cell_counts": censoring_counts,
        "unique_censoring_formula_count": len(formula_syntax),
        "unresolved_censoring_cells": len(unresolved),
        "ND_MDL_definition": "ND and <MDL form one category because both use MDL/sqrt(2) in the compiled dataset.",
        "MQL_definition": "Values below MQL use MQL/2 in the submitted dataset.",
        "unmeasured_definition": "Blank cells remain missing in every scheme.",
        "schemes": {name: vars(scheme) for name, scheme in SCHEMES.items()},
        "modeled_schemes": arguments.schemes,
        "analysis_mode": arguments.analysis_mode,
        "sensitivity_validation": arguments.validation,
        "seed": SEED,
        "grid_search_jobs": N_JOBS,
        "validation_design": {
            "outer_folds": N_OUTER_FOLDS,
            "inner_folds": N_INNER_FOLDS,
            "spatial_blocks": "Three equal-width longitude bins crossed with three latitude bins",
            "inner_groups": "Sampling_Group for grouped sensitivity comparisons",
            "grouped_target_treatment": "Outer-training 2.5th/97.5th percentile clipping before inner CV; outer-test targets unchanged",
            "legacy_target_treatment": "Original random-split reproduction includes whole-dataset clipping and is not an independent validation estimate",
            "endpoint": "GDR preprocessing sensitivity, separate from primary Magnitude-Balance validation",
        },
        "S0_cell_reproduction": bool(cell_audit["S0_matches_workbook"].all()),
        "S0_GDR_reproduction": bool(target_audit["S0_matches_submitted_GDR"].all()),
    }

    if arguments.analysis_mode != "prepare":
        features = primary.load_features(
            arguments.data,
            arguments.source_load,
            arguments.boundary,
            arguments.feature_table,
            arguments.routing_directory,
        )
        validate_feature_alignment(data, features)
        missing_covariates = [
            column for column in COVARIATES if column not in features.columns
        ]
        if missing_covariates:
            raise ValueError(f"Missing covariates: {', '.join(missing_covariates)}")
        features.to_csv(
            arguments.output_dir / "Censoring_Sensitivity_Feature_Table.csv",
            index=False,
        )
        manifest["feature_table"] = (
            str(arguments.feature_table.resolve()) if arguments.feature_table else None
        )
        manifest["source_load"] = (
            str(arguments.source_load.resolve()) if arguments.source_load else None
        )
        manifest["boundary"] = (
            str(arguments.boundary.resolve()) if not arguments.feature_table else None
        )

        if arguments.analysis_mode in {"all", "legacy"}:
            original, original_predictions, champions = fit_original_protocol(
                features, targets_by_scheme, arguments.schemes
            )
            paired, paired_predictions = fit_paired_fixed_sensitivity(
                features, targets_by_scheme, arguments.schemes
            )
            original.to_csv(
                arguments.output_dir / "Legacy_Original_Protocol_Performance.csv",
                index=False,
            )
            original_predictions.to_csv(
                arguments.output_dir / "Legacy_Original_Protocol_Predictions.csv",
                index=False,
            )
            champions.to_csv(
                arguments.output_dir / "Legacy_Original_Protocol_Champions.csv",
                index=False,
            )
            single_split_deltas(original).to_csv(
                arguments.output_dir / "Legacy_Original_Protocol_Deltas.csv",
                index=False,
            )
            paired.to_csv(
                arguments.output_dir / "Legacy_Paired_Fixed_Performance.csv",
                index=False,
            )
            paired_predictions.to_csv(
                arguments.output_dir / "Legacy_Paired_Fixed_Predictions.csv",
                index=False,
            )
            single_split_deltas(paired).to_csv(
                arguments.output_dir / "Legacy_Paired_Fixed_Deltas.csv", index=False
            )
            if arguments.reference_report is not None:
                historical_audit = compare_historical_report(
                    original, arguments.reference_report
                )
                historical_audit.to_csv(
                    arguments.output_dir / "S0_Historical_Model_Reproduction.csv",
                    index=False,
                )
                manifest["historical_report"] = str(
                    arguments.reference_report.resolve()
                )
                manifest["S0_historical_performance_reproduction"] = True

        if arguments.analysis_mode in {"all", "validation"}:
            per_fold, predictions = fit_sensitivity_validation(
                features, targets_by_scheme, arguments.schemes, arguments.validation
            )
            summary, deltas = summarize_performance(per_fold)
            per_fold.to_csv(
                arguments.output_dir / "Sensitivity_Performance_Per_Fold.csv",
                index=False,
            )
            summary.to_csv(
                arguments.output_dir / "Sensitivity_Performance_Summary.csv",
                index=False,
            )
            deltas.to_csv(
                arguments.output_dir / "Sensitivity_Performance_Deltas.csv", index=False
            )
            predictions.to_csv(
                arguments.output_dir / "Sensitivity_Out_of_Fold_Predictions.csv",
                index=False,
            )
            pooled = summarize_pooled_oof_performance(predictions)
            pooled.to_csv(
                arguments.output_dir / "Sensitivity_Pooled_OOF_Performance.csv",
                index=False,
            )

    primary.verify_inputs_unchanged(input_provenance)
    (arguments.output_dir / "Run_Manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
