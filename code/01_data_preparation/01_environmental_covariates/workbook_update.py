"""Apply environmental fields without rebuilding unrelated workbook content."""

from __future__ import annotations
from copy import deepcopy
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import uuid
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import (
    get_column_letter as letter,
    column_index_from_string as number,
)

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
TAG = lambda s: "{" + NS + "}" + s


def native_update(input_path, result_path, output_path):
    result = json.loads(Path(result_path).read_text(encoding="utf-8"))
    raw = Path(input_path).read_bytes()
    if hashlib.sha256(raw).hexdigest() != result["input_sha256"]:
        raise ValueError("Input workbook changed after extraction")
    before = load_workbook(io.BytesIO(raw), data_only=False)
    cache = load_workbook(io.BytesIO(raw), data_only=True)
    if len(before.worksheets) != 1:
        raise ValueError("Expected one observation sheet")
    sh = before.worksheets[0]
    heads = {c.value: c.column for c in sh[1] if c.value is not None}
    if not all(k in heads for k in ["TOC", "Temperature", "Temperature_Source"]):
        raise ValueError("Unexpected workbook schema")
    insertion = heads["TOC"] + 1 if "TOC_Source" not in heads else None
    move = lambda c: c + 1 if insertion and c >= insertion else c
    newheads = {h: move(c) for h, c in heads.items()}
    if insertion:
        newheads["TOC_Source"] = insertion

    def move_ref(text):
        return re.sub(
            r"(\$?)([A-Z]{1,3})(\$?\d+)",
            lambda m: m[1] + letter(move(number(m[2]))) + m[3],
            text,
        )

    changes = {}
    variable_changes = {}
    hydro = {}
    for record in result["records"]:
        r = record["excel_row"]
        f = record["variable"]
        val = record["value"]
        old = record["previous_value"]
        if str(sh.cell(r, heads["Number"]).value) != record["sample"]:
            raise ValueError("Record identity mismatch")
        actual = cache.worksheets[0].cell(r, heads[f]).value
        if actual != old:
            raise ValueError(f"Prior value mismatch: {f}, row {r}")
        if old is not None and val != old:
            raise ValueError("Original numerical value was modified")
        if val is not None:
            if not math.isfinite(val):
                raise ValueError("Nonfinite output")
            if old is None:
                changes[(f, r)] = val
                variable_changes[(f, r)] = val
        source = "Hydro_Source" if f.startswith("Hydro_") else f + "_Source"
        if source == "Hydro_Source":
            hydro.setdefault(r, []).append(record["source"])
        else:
            changes[(source, r)] = record["source"]
    for r, labels in hydro.items():
        present = set(labels) - {"Unfilled"}
        changes[("Hydro_Source", r)] = (
            next(iter(present))
            if len(present) == 1
            else "Unfilled" if not present else "Mixed Sources"
        )
    source_fields = {f for f, r in changes if f.endswith("_Source")}
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        members = z.infolist()
        parts = {p.filename: z.read(p.filename) for p in members}
    original_parts = parts.copy()
    xml = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
    styles = etree.fromstring(parts["xl/styles.xml"])
    fonts = styles.find(TAG("fonts"))
    xfs = styles.find(TAG("cellXfs"))
    nums = styles.find(TAG("numFmts"))
    oldcells = {c.get("r"): deepcopy(c) for c in xml.findall(".//" + TAG("c"))}
    template_numeric = oldcells["F2"]
    template_source = oldcells.get("G2", template_numeric)
    numeric_xf = deepcopy(xfs[int(template_numeric.get("s", "0"))])
    numeric_font = deepcopy(fonts[int(numeric_xf.get("fontId", "0"))])
    color = numeric_font.find(TAG("color"))
    if color is not None:
        numeric_font.remove(color)
    etree.SubElement(numeric_font, TAG("color"), rgb="FF000000")
    font_id = len(fonts)
    fonts.append(numeric_font)
    numeric_xf.set("fontId", str(font_id))
    numeric_xf.set("applyFont", "1")
    # Keep full numerical precision; formatting does not round stored values.
    numeric_style = len(xfs)
    xfs.append(numeric_xf)
    source_xf = deepcopy(xfs[int(template_source.get("s", "0"))])
    source_xf.set("fontId", str(font_id))
    source_xf.set("applyFont", "1")
    source_xf.set("numFmtId", "0")
    align = source_xf.find(TAG("alignment"))
    if align is None:
        align = etree.SubElement(source_xf, TAG("alignment"))
    align.set("horizontal", "center")
    align.set("vertical", "center")
    align.set("wrapText", "0")
    source_style = len(xfs)
    xfs.append(source_xf)
    source_header = deepcopy(xfs[int(oldcells["G1"].get("s", "0"))])
    header_style = len(xfs)
    xfs.append(source_header)

    def set_value(cell, value):
        for child in list(cell):
            if etree.QName(child).localname in ["f", "v", "is"]:
                cell.remove(child)
        cell.attrib.pop("t", None)
        if isinstance(value, str):
            cell.set("t", "inlineStr")
            etree.SubElement(etree.SubElement(cell, TAG("is")), TAG("t")).text = value
        else:
            etree.SubElement(cell, TAG("v")).text = format(value, ".17g")

    for row in xml.find(TAG("sheetData")):
        r = int(row.get("r"))
        cells = {}
        for cell in row.findall(TAG("c")):
            old_address = cell.get("r")
            old_col = number(re.match("[A-Z]+", old_address)[0])
            cell.set("r", letter(move(old_col)) + str(r))
            formula = cell.find(TAG("f"))
            if formula is not None:
                original = sh[old_address].value
                if not isinstance(original, str) or not original.startswith("="):
                    raise ValueError("Unresolved formula")
                formula.attrib.clear()
                formula.text = move_ref(original[1:])
            cells[move(old_col)] = cell
        if r == 1 and insertion:
            cell = etree.Element(
                TAG("c"), r=letter(insertion) + "1", s=str(header_style)
            )
            set_value(cell, "TOC_Source")
            cells[insertion] = cell
        for (f, rr), value in changes.items():
            if rr != r:
                continue
            col = newheads[f]
            cell = cells.get(col)
            if cell is None:
                cell = etree.Element(TAG("c"), r=letter(col) + str(r))
                cells[col] = cell
            set_value(cell, value)
            cell.set("s", str(source_style if f in source_fields else numeric_style))
        for cell in row.findall(TAG("c")):
            row.remove(cell)
        for _, cell in sorted(cells.items()):
            row.append(cell)
        if row.get("spans"):
            row.set("spans", f"1:{max(newheads.values())}")
    # Shift native range references together with the inserted source column.
    for element in xml.iter():
        if element.tag in [TAG("c"), TAG("f")]:
            continue
        for attr in ("ref", "sqref", "activeCell", "topLeftCell"):
            if element.get(attr):
                element.set(attr, move_ref(element.get(attr)))
    cols = xml.find(TAG("cols"))
    if cols is None:
        raise ValueError("Expected original column definitions")
    originals = list(cols)
    cols.clear()
    for field, c in sorted(newheads.items(), key=lambda p: p[1]):
        old = heads.get(field, heads["Temperature_Source"])
        original = next(
            (e for e in originals if int(e.get("min")) <= old <= int(e.get("max"))),
            None,
        )
        element = (
            deepcopy(original)
            if original is not None
            else etree.Element(TAG("col"), width="18", customWidth="1")
        )
        element.set("min", str(c))
        element.set("max", str(c))
        if field in source_fields:
            element.set("width", "32")
            element.set("customWidth", "1")
        elif field in {f for f, r in variable_changes}:
            element.set("width", str(max(24, float(element.get("width", "18")))))
            element.set("customWidth", "1")
        cols.append(element)
    fonts.set("count", str(len(fonts)))
    xfs.set("count", str(len(xfs)))
    parts["xl/styles.xml"] = etree.tostring(
        styles, encoding="utf-8", xml_declaration=True
    )
    parts["xl/worksheets/sheet1.xml"] = etree.tostring(
        xml, encoding="utf-8", xml_declaration=True
    )
    if insertion and "xl/calcChain.xml" in parts:
        chain = etree.fromstring(parts["xl/calcChain.xml"])
        for cell in chain:
            if cell.get("r"):
                cell.set("r", move_ref(cell.get("r")))
        parts["xl/calcChain.xml"] = etree.tostring(
            chain, encoding="utf-8", xml_declaration=True
        )
    if insertion:
        wx = etree.fromstring(parts["xl/workbook.xml"])
        for element in wx.findall(".//" + TAG("definedName")):
            if element.text:
                element.text = move_ref(element.text)
        parts["xl/workbook.xml"] = etree.tostring(
            wx, encoding="utf-8", xml_declaration=True
        )
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as z:
        for item in members:
            z.writestr(item, parts[item.filename])
    Path(output_path).write_bytes(stream.getvalue())
    after = load_workbook(output_path, data_only=False)
    ac = load_workbook(output_path, data_only=True)
    ash = after.worksheets[0]
    ach = ac.worksheets[0]
    checked = 0
    for f, oldcol in heads.items():
        newcol = newheads[f]
        for r in range(1, sh.max_row + 1):
            oldcell = sh.cell(r, oldcol)
            newcell = ash.cell(r, newcol)
            key = (f, r)
            if key in changes:
                expected = changes[key]
                if isinstance(expected, float):
                    assert math.isclose(
                        newcell.value, expected, rel_tol=1e-14, abs_tol=1e-14
                    )
                else:
                    assert newcell.value == expected
            else:
                expected = (
                    move_ref(oldcell.value)
                    if oldcell.data_type == "f"
                    else oldcell.value
                )
                assert newcell.value == expected, (f, r, "unexpected value")
                assert (
                    ach.cell(r, newcol).value
                    == cache.worksheets[0].cell(r, oldcol).value
                ), (f, r, "unexpected cache")
                assert oldcell._style == newcell._style, (f, r, "unexpected style")
                checked += 1
    for f, r in changes:
        c = ash.cell(r, newheads[f])
        assert c.font.color.type == "rgb" and c.font.color.rgb[-6:] == "000000"
    allowed = {
        "xl/styles.xml",
        "xl/worksheets/sheet1.xml",
        "xl/calcChain.xml",
        "xl/workbook.xml",
    }
    assert all(parts[k] == v for k, v in original_parts.items() if k not in allowed)
    assert before.sheetnames == after.sheetnames
    for book in (before, cache, after, ac):
        book.close()
    return {
        "input_sha256": result["input_sha256"],
        "output_sha256": hashlib.sha256(stream.getvalue()).hexdigest(),
        "unchanged_cells_verified": checked,
        "environment_values_added": len(variable_changes),
        "source_cells_normalized": len(changes) - len(variable_changes),
        "columns": newheads,
        "native_parts_preserved": True,
    }


def publish_result(config):
    """Back up and publish a verified candidate, with a separate provenance file."""
    from environment_core import dump_json, sha256

    source = Path(config["workbook"]).resolve()
    output = Path(config["output_directory"]).resolve()
    result_path = output / "environment_results.json"
    result = json.loads(result_path.read_text(encoding="utf-8"))
    if (source.parent / ("~$" + source.name)).exists():
        raise RuntimeError("Close the workbook in Excel before writeback")
    if sha256(source) != result["input_sha256"]:
        raise RuntimeError("Workbook changed since extraction; writeback cancelled")
    backup_dir = output / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    original = source.read_bytes()
    backup = backup_dir / (
        "DATA_before_environment_" + result["input_sha256"][:12] + ".xlsx"
    )
    if backup.exists():
        if sha256(backup) != result["input_sha256"]:
            raise RuntimeError("Existing backup does not match original")
    else:
        backup.write_bytes(original)
    candidate = output / "verified_DATA.xlsx"
    report = native_update(source, result_path, candidate)
    record_path = source.parent / "Environment_preparation_record.json"
    old_record = record_path.read_bytes() if record_path.exists() else None
    if old_record is not None:
        (
            backup_dir
            / (
                "Environment_record_"
                + hashlib.sha256(old_record).hexdigest()[:12]
                + ".json"
            )
        ).write_bytes(old_record)
    record = {
        "scope": "01 environmental covariates only",
        "workbook": str(source),
        "backup": str(backup),
        "verification": report,
        "summary": json.loads(
            (output / "environment_summary.json").read_text(encoding="utf-8")
        ),
        "execution_checkpoint": json.loads(
            (output / "checkpoint.json").read_text(encoding="utf-8")
        ),
        "publication_code": {
            p.name: sha256(p) for p in sorted(Path(__file__).parent.glob("*.py"))
        },
        "source_label_policy": "Short source labels in DATA; previous labels and per-cell extraction details retained here. Original Data includes previously approved source-derived observations; see Dataset_preparation_record.json for their earlier audit trail.",
        "format_policy": "New environmental values and source labels are black. Existing analytical-limit formulas and all unrelated cell styles are preserved.",
        "validation_note": "Observation-dependent imputations are marked requires_fold_refit in per-cell details. They must be evaluated within the subsequent modeling-validation pipeline; filling the workbook does not itself establish model validity.",
        "extraction": result,
    }
    token = uuid.uuid4().hex
    staged_data = source.parent / (".environment_" + token + ".xlsx")
    staged_record = source.parent / (".environment_" + token + ".json")
    data_replaced = False
    record_replaced = False
    try:
        staged_data.write_bytes(candidate.read_bytes())
        dump_json(staged_record, record)
        if sha256(source) != result["input_sha256"]:
            raise RuntimeError("Workbook changed before publication")
        os.replace(staged_data, source)
        data_replaced = True
        os.replace(staged_record, record_path)
        record_replaced = True
        assert sha256(source) == report["output_sha256"]
    except Exception:
        if data_replaced:
            staged_data.write_bytes(original)
            os.replace(staged_data, source)
        if record_replaced:
            if old_record is None:
                record_path.unlink()
            else:
                staged_record.write_bytes(old_record)
                os.replace(staged_record, record_path)
        raise
    finally:
        for temporary in (staged_data, staged_record):
            if temporary.exists():
                temporary.unlink()
    receipt = dict(
        report, workbook=str(source), provenance=str(record_path), backup=str(backup)
    )
    dump_json(output / "publication_receipt.json", receipt)
    return receipt
